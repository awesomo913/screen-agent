"""
toolkit_34_word_excel.py
Read and write Microsoft Word (.docx) and Excel (.xlsx) files
via python-docx and openpyxl (soft-import).
"""
from __future__ import annotations
import os
from pathlib import Path
from typing import Any, Dict, List

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

def check_libraries() -> Dict[str, Any]:
    return {"success": True, "data": {"python-docx": HAS_DOCX, "openpyxl": HAS_OPENPYXL}, "error": None}

# --- WORD ---

def read_docx(path: str) -> Dict[str, Any]:
    try:
        if not HAS_DOCX:
            return {"success": False, "data": None, "error": "python-docx not installed (pip install python-docx)"}
        doc = DocxDocument(path)
        paras = [p.text for p in doc.paragraphs]
        full_text = "\n".join(paras)
        return {"success": True, "data": {"paragraphs": paras, "full_text": full_text, "paragraph_count": len(paras)}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def get_docx_info(path: str) -> Dict[str, Any]:
    try:
        if not HAS_DOCX:
            return {"success": False, "data": None, "error": "python-docx not installed"}
        doc = DocxDocument(path)
        core = doc.core_properties
        return {"success": True, "data": {
            "title": core.title,
            "author": core.author,
            "created": str(core.created),
            "modified": str(core.modified),
            "subject": core.subject,
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "sections": len(doc.sections),
            "size_kb": round(os.path.getsize(path) / 1024, 1)
        }, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def create_docx(output_path: str, title: str = "", content: str = "") -> Dict[str, Any]:
    try:
        if not HAS_DOCX:
            return {"success": False, "data": None, "error": "python-docx not installed"}
        doc = DocxDocument()
        if title:
            doc.add_heading(title, level=1)
        if content:
            for para in content.split("\n"):
                doc.add_paragraph(para)
        doc.save(output_path)
        return {"success": True, "data": {"saved": output_path}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def append_to_docx(path: str, text: str, heading: str = "") -> Dict[str, Any]:
    try:
        if not HAS_DOCX:
            return {"success": False, "data": None, "error": "python-docx not installed"}
        doc = DocxDocument(path)
        if heading:
            doc.add_heading(heading, level=2)
        doc.add_paragraph(text)
        doc.save(path)
        return {"success": True, "data": "Text appended", "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def docx_to_text(path: str, output_txt: str) -> Dict[str, Any]:
    try:
        if not HAS_DOCX:
            return {"success": False, "data": None, "error": "python-docx not installed"}
        doc = DocxDocument(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        with open(output_txt, "w", encoding="utf-8") as f:
            f.write(text)
        return {"success": True, "data": {"saved": output_txt, "chars": len(text)}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def read_docx_tables(path: str) -> Dict[str, Any]:
    try:
        if not HAS_DOCX:
            return {"success": False, "data": None, "error": "python-docx not installed"}
        doc = DocxDocument(path)
        tables = []
        for t in doc.tables:
            rows = [[cell.text for cell in row.cells] for row in t.rows]
            tables.append(rows)
        return {"success": True, "data": {"table_count": len(tables), "tables": tables}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

# --- EXCEL ---

def read_xlsx(path: str, sheet_name: str = "", max_rows: int = 1000) -> Dict[str, Any]:
    try:
        if not HAS_OPENPYXL:
            return {"success": False, "data": None, "error": "openpyxl not installed (pip install openpyxl)"}
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append(list(row))
        wb.close()
        return {"success": True, "data": {"sheet": ws.title, "rows": rows, "row_count": len(rows)}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def get_xlsx_info(path: str) -> Dict[str, Any]:
    try:
        if not HAS_OPENPYXL:
            return {"success": False, "data": None, "error": "openpyxl not installed"}
        wb = openpyxl.load_workbook(path, read_only=True)
        info = {
            "sheets": wb.sheetnames,
            "sheet_count": len(wb.sheetnames),
            "size_kb": round(os.path.getsize(path) / 1024, 1)
        }
        wb.close()
        return {"success": True, "data": info, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def create_xlsx(output_path: str, data: list, sheet_name: str = "Sheet1") -> Dict[str, Any]:
    """Create an Excel file from a list of lists (rows)."""
    try:
        if not HAS_OPENPYXL:
            return {"success": False, "data": None, "error": "openpyxl not installed"}
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in data:
            ws.append(row)
        wb.save(output_path)
        return {"success": True, "data": {"saved": output_path, "rows": len(data)}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def append_xlsx_rows(path: str, rows: list, sheet_name: str = "") -> Dict[str, Any]:
    try:
        if not HAS_OPENPYXL:
            return {"success": False, "data": None, "error": "openpyxl not installed"}
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        for row in rows:
            ws.append(row)
        wb.save(path)
        return {"success": True, "data": {"appended_rows": len(rows)}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def xlsx_to_csv(xlsx_path: str, csv_path: str, sheet_name: str = "") -> Dict[str, Any]:
    try:
        if not HAS_OPENPYXL:
            return {"success": False, "data": None, "error": "openpyxl not installed"}
        import csv
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow([str(c) if c is not None else "" for c in row])
        wb.close()
        return {"success": True, "data": {"csv_saved": csv_path}, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}

def get_xlsx_sheet_dimensions(path: str, sheet_name: str = "") -> Dict[str, Any]:
    try:
        if not HAS_OPENPYXL:
            return {"success": False, "data": None, "error": "openpyxl not installed"}
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        return {"success": True, "data": {
            "sheet": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column
        }, "error": None}
    except Exception as e:
        return {"success": False, "data": None, "error": str(e)}
